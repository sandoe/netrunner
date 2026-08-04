from fastapi import APIRouter, HTTPException, Response
from pydantic import BaseModel
from typing import Dict, Any, List
import uuid
import time
import io

try:
    from docx import Document
except ImportError:
    Document = None

from backend.routers.nodes import load_nodes
from backend.core.db import (
    load_alerts_db,
    load_links_db,
    save_report_db,
    load_reports_db,
    get_report_db,
    delete_report_db,
)

router = APIRouter(prefix="/reports", tags=["reports"])


@router.get("", response_model=Dict[str, Any])
async def get_pentest_report():
    """Aggregates data for pentest reporting."""
    nodes = await load_nodes()
    links = await load_links_db()
    alerts = await load_alerts_db()

    return {"nodes": nodes, "links": links, "alerts": alerts}


@router.get("/summary", response_model=Dict[str, Any])
async def get_report_summary(timerange_hours: int = 24):
    nodes = await load_nodes()
    alerts = await load_alerts_db()

    # Filter alerts by time range if needed, for now just basic count
    now = time.time()
    recent_alerts = [
        a for a in alerts if now - a["created_at"] <= timerange_hours * 3600
    ]

    severity_counts = {"low": 0, "medium": 0, "high": 0, "critical": 0}
    status_counts = {"new": 0, "open": 0, "closed": 0, "false_positive": 0}
    for a in recent_alerts:
        severity_counts[a.get("severity", "low")] += 1
        status_counts[a.get("status", "new")] += 1

    return {
        "generated_at": now,
        "timerange_hours": timerange_hours,
        "total_active_nodes": len(nodes),
        "alerts_summary": {
            "total": len(recent_alerts),
            "by_severity": severity_counts,
            "by_status": status_counts,
        },
        "threats_summary": {"total_events": 0},  # Placeholder, would load threat events
    }


class SaveReportRequest(BaseModel):
    timerange_hours: int
    summary_json: Dict[str, Any]
    markdown_content: str


@router.post("/save")
async def save_pentest_report(req: SaveReportRequest):
    report_id = str(uuid.uuid4())
    report_data = {
        "id": report_id,
        "timerange_hours": req.timerange_hours,
        "summary_json": req.summary_json,
        "markdown_content": req.markdown_content,
        "created_at": time.time(),
    }
    await save_report_db(report_data)
    return {"status": "success", "report_id": report_id}


@router.get("/history")
async def get_report_history():
    reports = await load_reports_db()
    return reports


@router.get("/{report_id}/export")
async def export_report(report_id: str, format: str = "md"):
    report = await get_report_db(report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")

    md_content = report["markdown_content"]

    if format == "md":
        return Response(
            content=md_content,
            media_type="text/markdown",
            headers={
                "Content-Disposition": f'attachment; filename="report_{report_id}.md"'
            },
        )
    elif format == "latex":
        # A simple fallback for LaTeX representation
        latex_content = f"\\documentclass{{article}}\n\\usepackage[utf8]{{inputenc}}\n\\title{{Pentest Report}}\n\\begin{{document}}\n\\maketitle\n\n"
        for line in md_content.split("\n"):
            if line.startswith("# "):
                title = line[2:].replace("_", "\\_")
                latex_content += f"\\section{{{title}}}\n"
            elif line.startswith("## "):
                title = line[3:].replace("_", "\\_")
                latex_content += f"\\subsection{{{title}}}\n"
            elif line.startswith("- "):
                content = line[2:].replace("_", "\\_")
                latex_content += f"\\textbullet\\ {content}\\\\\n"
            else:
                content = line.replace("_", "\\_")
                latex_content += f"{content}\n\n"
        latex_content += "\\end{document}\n"

        return Response(
            content=latex_content,
            media_type="application/x-latex",
            headers={
                "Content-Disposition": f'attachment; filename="report_{report_id}.tex"'
            },
        )
    elif format == "docx":
        if Document is None:
            raise HTTPException(
                status_code=500, detail="python-docx library not installed"
            )

        doc = Document()
        doc.add_heading("Pentest Report", 0)
        # simplistic conversion
        for line in md_content.split("\n"):
            line_clean = line.strip()
            if not line_clean:
                continue
            if line_clean.startswith("# "):
                doc.add_heading(line_clean[2:], level=1)
            elif line_clean.startswith("## "):
                doc.add_heading(line_clean[3:], level=2)
            elif line_clean.startswith("### "):
                doc.add_heading(line_clean[4:], level=3)
            elif line_clean.startswith("- "):
                doc.add_paragraph(line_clean[2:], style="List Bullet")
            else:
                doc.add_paragraph(line_clean)

        f = io.BytesIO()
        doc.save(f)
        f.seek(0)
        return Response(
            content=f.read(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f'attachment; filename="report_{report_id}.docx"'
            },
        )
    else:
        raise HTTPException(
            status_code=400, detail="Invalid format. Supported: md, latex, docx"
        )


@router.delete("/{report_id}")
async def delete_report(report_id: str):
    report = await get_report_db(report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")

    await delete_report_db(report_id)
    return {"status": "success", "message": f"Report {report_id} deleted"}
